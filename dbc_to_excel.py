#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DBC (CAN) -> Excel (.xlsx)

Application Windows (Tkinter) qui analyse entièrement un fichier DBC et exporte
toutes ses informations dans UN SEUL fichier Excel, UN SEUL onglet, UN SEUL
tableau structuré (openpyxl Table) avec UNE LIGNE PAR SIGNAL.

Dépendance unique :  pip install openpyxl
Lancement GUI     :  python dbc_to_excel.py
Lancement console :  python dbc_to_excel.py entree.dbc [sortie.xlsx]
"""
from __future__ import annotations

import datetime
import os
import queue
import re
import sys
import threading
from dataclasses import dataclass, field
from typing import Callable, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# =============================================================================
#  MODÈLE DE DONNÉES
# =============================================================================


@dataclass
class AttrDef:
    obj: str                      # '' (global), 'BU_', 'BO_', 'SG_', 'EV_'
    name: str
    kind: str                     # INT, HEX, FLOAT, STRING, ENUM
    enum: list = field(default_factory=list)
    minimum: Optional[float] = None
    maximum: Optional[float] = None


@dataclass
class Signal:
    name: str
    mux_indicator: str            # '', 'M', 'm3', 'm3M'
    start_bit: int
    length: int
    byte_order: int               # 1 = Intel (little endian), 0 = Motorola (big endian)
    signed: bool
    factor: float
    offset: float
    minimum: Optional[float]
    maximum: Optional[float]
    unit: str
    receivers: list
    value_type: int = 0           # 0 = entier, 1 = float 32, 2 = float 64 (SIG_VALTYPE_)
    comment: str = ""
    values: dict = field(default_factory=dict)      # valeur brute -> texte
    attrs: dict = field(default_factory=dict)       # attributs explicites (résolus)
    mux_switch: Optional[str] = None                # nom du signal multiplexeur
    mux_ranges: list = field(default_factory=list)  # [(min, max), ...]


@dataclass
class Message:
    frame_id: int                 # valeur brute du DBC (bit 31 = trame étendue)
    name: str
    dlc: int
    sender: str
    signals: list = field(default_factory=list)
    extra_senders: list = field(default_factory=list)   # BO_TX_BU_
    comment: str = ""
    attrs: dict = field(default_factory=dict)


@dataclass
class DbcDatabase:
    source: str = ""
    version: str = ""
    comment: str = ""
    nodes: dict = field(default_factory=dict)            # nom -> commentaire
    messages: list = field(default_factory=list)
    attr_defs: dict = field(default_factory=dict)        # (obj, nom) -> AttrDef
    attr_defaults: dict = field(default_factory=dict)    # (obj, nom) -> valeur résolue
    global_attrs: dict = field(default_factory=dict)
    value_tables: dict = field(default_factory=dict)
    used_attrs: list = field(default_factory=list)       # [(obj, nom)] réellement renseignés


# =============================================================================
#  PARSEUR DBC
# =============================================================================

_OBJ = ("BU_", "BO_", "SG_", "EV_")
_TOKEN_RE = re.compile(r'"(?:[^"\\]|\\.)*"|[^\s;,"]+')
_SG_RE = re.compile(
    r'^\s*SG_\s+(\w+)\s*(M|m\d+M?)?\s*:\s*(\d+)\|(\d+)@([01])([+-])\s*'
    r'\(\s*([^,\s]+)\s*,\s*([^)\s]+)\s*\)\s*'
    r'\[\s*([^|\s]*)\s*\|\s*([^\]\s]*)\s*\]\s*'
    r'"((?:[^"\\]|\\.)*)"\s*(.*?)\s*$'
)
_BO_RE = re.compile(r'^BO_\s+(\d+)\s+(\w+)\s*:\s*(\d+)\s+(\w+)')
_STMT_KEYS = {
    "CM_", "BA_DEF_", "BA_DEF_DEF_", "BA_", "VAL_", "VAL_TABLE_", "SIG_VALTYPE_",
    "BO_TX_BU_", "SG_MUL_VAL_", "BA_DEF_REL_", "BA_DEF_DEF_REL_", "BA_REL_",
    "SIG_GROUP_", "EV_", "ENVVAR_DATA_", "SGTYPE_", "SGTYPE_VAL_", "BA_DEF_SGTYPE_",
    "BA_SGTYPE_", "SIG_TYPE_REF_", "SIGTYPE_VALTYPE_", "CAT_DEF_", "CAT_", "FILTER",
    "BU_SG_REL_", "BU_EV_REL_", "BU_BO_REL_",
}


def _read_text(path: str) -> str:
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _unq(tok: str) -> str:
    if len(tok) >= 2 and tok[0] == '"' and tok[-1] == '"':
        tok = tok[1:-1]
    return re.sub(r'\\(["\\])', r"\1", tok)


def _num(tok):
    if tok is None or tok == "":
        return None
    try:
        return int(tok)
    except ValueError:
        pass
    try:
        f = float(tok)
    except ValueError:
        return None
    if f == int(f) and abs(f) < 1e15:
        return int(f)
    return f


def _complete(buf: str) -> bool:
    quotes = len(re.findall(r'(?<!\\)"', buf))
    return quotes % 2 == 0 and buf.rstrip().endswith(";")


def _tokens(stmt: str) -> list:
    return [t for t in _TOKEN_RE.findall(stmt) if t != ":"]


def _pairs(tokens: list) -> dict:
    out = {}
    for i in range(0, len(tokens) - 1, 2):
        n = _num(tokens[i])
        if n is not None:
            out[int(n)] = _unq(tokens[i + 1])
    return out


def _attr_value(tok: str):
    return _unq(tok) if tok.startswith('"') else _num(tok)


def parse_dbc(path: str) -> DbcDatabase:
    db = DbcDatabase(source=os.path.basename(path))
    text = _read_text(path).replace("\r\n", "\n").replace("\r", "\n")

    messages: dict = {}                 # id -> Message
    cur: Optional[Message] = None
    sig_comments, msg_comments, node_comments = {}, {}, {}
    sig_values, raw_attrs = {}, {}      # raw_attrs[(obj, key)][nom] = valeur brute
    val_types, mux_ext, extra_tx = {}, {}, {}
    defaults_raw = {}

    def handle(stmt: str):
        t = _tokens(stmt)
        if not t:
            return
        k = t[0]
        if k == "CM_":
            if len(t) >= 2 and t[1] in _OBJ:
                if t[1] == "BU_" and len(t) >= 4:
                    node_comments[t[2]] = _unq(t[3])
                elif t[1] == "BO_" and len(t) >= 4:
                    msg_comments[int(t[2])] = _unq(t[3])
                elif t[1] == "SG_" and len(t) >= 5:
                    sig_comments[(int(t[2]), t[3])] = _unq(t[4])
            elif len(t) >= 2:
                db.comment = _unq(t[1])
        elif k == "BA_DEF_":
            if t[1] in _OBJ:
                obj, name, kind, params = t[1], _unq(t[2]), t[3], t[4:]
            else:
                obj, name, kind, params = "", _unq(t[1]), t[2], t[3:]
            d = AttrDef(obj, name, kind)
            if kind == "ENUM":
                d.enum = [_unq(p) for p in params]
            elif kind in ("INT", "HEX", "FLOAT") and len(params) >= 2:
                d.minimum, d.maximum = _num(params[0]), _num(params[1])
            db.attr_defs[(obj, name)] = d
        elif k == "BA_DEF_DEF_":
            if len(t) >= 3:
                defaults_raw[_unq(t[1])] = _attr_value(t[2])
        elif k == "BA_":
            name = _unq(t[1])
            if len(t) > 2 and t[2] in _OBJ:
                obj = t[2]
                if obj in ("BU_", "EV_") and len(t) >= 5:
                    raw_attrs.setdefault((obj, t[3]), {})[name] = _attr_value(t[4])
                elif obj == "BO_" and len(t) >= 5:
                    raw_attrs.setdefault((obj, int(t[3])), {})[name] = _attr_value(t[4])
                elif obj == "SG_" and len(t) >= 6:
                    raw_attrs.setdefault((obj, (int(t[3]), t[4])), {})[name] = _attr_value(t[5])
            elif len(t) >= 3:
                raw_attrs.setdefault(("", ""), {})[name] = _attr_value(t[2])
        elif k == "VAL_":
            if len(t) >= 3 and re.fullmatch(r"\d+", t[1]):
                sig_values[(int(t[1]), t[2])] = _pairs(t[3:])
        elif k == "VAL_TABLE_":
            if len(t) >= 2:
                db.value_tables[t[1]] = _pairs(t[2:])
        elif k == "SIG_VALTYPE_":
            if len(t) >= 4:
                val_types[(int(t[1]), t[2])] = int(t[3])
        elif k == "BO_TX_BU_":
            if len(t) >= 3:
                extra_tx[int(t[1])] = t[2:]
        elif k == "SG_MUL_VAL_":
            if len(t) >= 4:
                ranges = [(int(a), int(b)) for a, b in
                          (re.fullmatch(r"(\d+)-(\d+)", x).groups()
                           for x in t[4:] if re.fullmatch(r"\d+-\d+", x))]
                mux_ext[(int(t[1]), t[2])] = (t[3], ranges)

    buf = None
    for raw in text.split("\n"):
        if buf is not None:
            buf += "\n" + raw
            if _complete(buf):
                handle(buf)
                buf = None
            continue
        if not raw.strip():
            continue
        first = raw.split(None, 1)[0]
        if raw[0] in " \t" or first == "SG_":
            m = _SG_RE.match(raw)
            if m and cur is not None:
                (name, mux, sb, ln, bo, sg, fac, off, mn, mx, unit, rcv) = m.groups()
                cur.signals.append(Signal(
                    name=name, mux_indicator=mux or "", start_bit=int(sb), length=int(ln),
                    byte_order=int(bo), signed=(sg == "-"),
                    factor=_num(fac) if _num(fac) is not None else 1,
                    offset=_num(off) if _num(off) is not None else 0,
                    minimum=_num(mn), maximum=_num(mx), unit=_unq('"' + unit + '"'),
                    receivers=[r for r in re.split(r"[,\s]+", rcv) if r]))
            continue
        if first == "BO_":
            m = _BO_RE.match(raw)
            if m:
                cur = Message(int(m.group(1)), m.group(2), int(m.group(3)), m.group(4))
                messages[cur.frame_id] = cur
                db.messages.append(cur)
            continue
        mb = re.match(r"^BU_\s*:(.*)$", raw)
        if mb:
            for n in mb.group(1).split():
                db.nodes.setdefault(n, "")
            continue
        if first == "VERSION":
            mv = re.match(r'^VERSION\s+"((?:[^"\\]|\\.)*)"', raw)
            db.version = _unq(mv.group(1)) if mv else ""
            continue
        if first in _STMT_KEYS:
            cur = None
            if _complete(raw):
                handle(raw)
            else:
                buf = raw

    # ---- résolution des attributs (index ENUM -> texte) ---------------------
    def resolve(obj, name, val):
        d = db.attr_defs.get((obj, name))
        if d and d.kind == "ENUM" and isinstance(val, (int, float)) and not isinstance(val, bool):
            idx = int(val)
            if 0 <= idx < len(d.enum):
                return d.enum[idx]
        return val

    for (obj, name), d in db.attr_defs.items():
        if name in defaults_raw:
            db.attr_defaults[(obj, name)] = resolve(obj, name, defaults_raw[name])
    for name, val in raw_attrs.get(("", ""), {}).items():
        db.global_attrs[name] = resolve("", name, val)

    used = set()
    for msg in db.messages:
        msg.comment = msg_comments.get(msg.frame_id, "")
        msg.extra_senders = extra_tx.get(msg.frame_id, [])
        for n, v in raw_attrs.get(("BO_", msg.frame_id), {}).items():
            msg.attrs[n] = resolve("BO_", n, v)
            used.add(("BO_", n))
        roots = [s.name for s in msg.signals if s.mux_indicator == "M"]
        for s in msg.signals:
            key = (msg.frame_id, s.name)
            s.comment = sig_comments.get(key, "")
            s.values = sig_values.get(key, {})
            s.value_type = val_types.get(key, 0)
            for n, v in raw_attrs.get(("SG_", key), {}).items():
                s.attrs[n] = resolve("SG_", n, v)
                used.add(("SG_", n))
            if s.mux_indicator.startswith("m"):
                num = int(re.match(r"m(\d+)", s.mux_indicator).group(1))
                if key in mux_ext:
                    s.mux_switch, s.mux_ranges = mux_ext[key][0], mux_ext[key][1] or [(num, num)]
                else:
                    s.mux_switch = roots[0] if roots else None
                    s.mux_ranges = [(num, num)]
    for n, c in node_comments.items():
        db.nodes[n] = c
    db.used_attrs = [k for k in db.attr_defs if k in used]
    return db


# =============================================================================
#  EXPORT EXCEL
# =============================================================================

_FONT = "Arial"
_HDR_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
_THIN = Side(style="thin", color="A6A6A6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MAX_CELL = 32000


@dataclass
class Col:
    header: str
    fn: Callable
    fmt: Optional[str] = None
    wrap: bool = False
    max_width: int = 45


def _can_id(m: Message):
    fid = m.frame_id
    ext = bool(fid & 0x80000000) or fid > 0x7FF
    if ext:
        return True, f"0x{fid & 0x1FFFFFFF:08X}"
    return False, f"0x{fid & 0x7FF:03X}"


def _ranges_text(ranges) -> str:
    return ", ".join(str(a) if a == b else f"{a}-{b}" for a, b in ranges)


def _on_sig(f):
    return lambda m, s: f(m, s) if s is not None else None


def _mux_role(s: Signal) -> str:
    ind = s.mux_indicator
    if ind == "M":
        return "Multiplexeur"
    if ind.startswith("m") and ind.endswith("M"):
        return "Multiplexé + multiplexeur"
    if ind.startswith("m"):
        return "Multiplexé"
    return "Non multiplexé"


def _mux_value(s: Signal):
    if not s.mux_ranges:
        return None
    if len(s.mux_ranges) == 1 and s.mux_ranges[0][0] == s.mux_ranges[0][1]:
        return s.mux_ranges[0][0]
    return _ranges_text(s.mux_ranges)


def _mux_cond(s: Signal):
    if not s.mux_ranges:
        return None
    sw = s.mux_switch or "?"
    if len(s.mux_ranges) == 1 and s.mux_ranges[0][0] == s.mux_ranges[0][1]:
        return f"{sw} = {s.mux_ranges[0][0]}"
    return f"{sw} dans [{_ranges_text(s.mux_ranges)}]"


def _value_type(s: Signal) -> str:
    if s.value_type == 1:
        return "Float 32 bits"
    if s.value_type == 2:
        return "Float 64 bits"
    return "Signé" if s.signed else "Non signé"


def _build_columns(db: DbcDatabase, include_attrs: bool) -> list:
    c = [
        Col("Message", lambda m, s: m.name, max_width=50),
        Col("Commentaire message", lambda m, s: m.comment, wrap=True, max_width=60),
        Col("ID (décimal DBC)", lambda m, s: m.frame_id, "0"),
        Col("ID CAN (hex)", lambda m, s: _can_id(m)[1]),
        Col("Format de trame", lambda m, s: "Étendu (29 bits)" if _can_id(m)[0] else "Standard (11 bits)"),
        Col("DLC (octets)", lambda m, s: m.dlc, "0"),
        Col("Émetteur", lambda m, s: m.sender),
        Col("Émetteurs additionnels", lambda m, s: ", ".join(m.extra_senders), wrap=True),
        Col("Message multiplexé",
            lambda m, s: "Oui" if any(x.mux_indicator for x in m.signals) else "Non"),
        Col("Signal", _on_sig(lambda m, s: s.name), max_width=50),
        Col("Commentaire signal", _on_sig(lambda m, s: s.comment), wrap=True, max_width=60),
        Col("Bit de départ", _on_sig(lambda m, s: s.start_bit), "0"),
        Col("Longueur (bits)", _on_sig(lambda m, s: s.length), "0"),
        Col("Ordre des octets",
            _on_sig(lambda m, s: "Intel (little-endian)" if s.byte_order == 1 else "Motorola (big-endian)")),
        Col("Type de valeur", _on_sig(lambda m, s: _value_type(s))),
        Col("Facteur", _on_sig(lambda m, s: s.factor), "General"),
        Col("Offset", _on_sig(lambda m, s: s.offset), "General"),
        Col("Min", _on_sig(lambda m, s: s.minimum), "General"),
        Col("Max", _on_sig(lambda m, s: s.maximum), "General"),
        Col("Unité", _on_sig(lambda m, s: s.unit)),
        Col("Récepteurs", _on_sig(lambda m, s: ", ".join(s.receivers)), wrap=True),
        Col("Rôle multiplexage", _on_sig(lambda m, s: _mux_role(s))),
        Col("Signal multiplexeur", _on_sig(lambda m, s: s.mux_switch)),
        Col("Valeur(s) du multiplexeur", _on_sig(lambda m, s: _mux_value(s)), "0"),
        Col("Condition de présence", _on_sig(lambda m, s: _mux_cond(s)), wrap=True),
        Col("Table de valeurs",
            _on_sig(lambda m, s: "\n".join(f"{k} = {v}" for k, v in s.values.items())),
            wrap=True, max_width=50),

    
        Col("Commentaire nœud émetteur", lambda m, s: db.nodes.get(m.sender, ""),
            wrap=True, max_width=50),
    ]
    if include_attrs:
        for obj, name in db.used_attrs:
            dflt = db.attr_defaults.get((obj, name))
            if obj == "BO_":
                c.append(Col(f"Msg.{name}",
                             lambda m, s, n=name, d=dflt: m.attrs.get(n, d), "General"))
            elif obj == "SG_":
                c.append(Col(f"Sig.{name}",
                             _on_sig(lambda m, s, n=name, d=dflt: s.attrs.get(n, d)), "General"))
    # en-têtes strictement uniques (exigence des tableaux Excel)
    seen = {}
    for col in c:
        h = col.header
        seen[h] = seen.get(h, 0) + 1
        if seen[h] > 1:
            col.header = f"{h} ({seen[h]})"
    return c


def _clean(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > _MAX_CELL:
            v = v[:_MAX_CELL] + "… [tronqué]"
        return v if v != "" else None
    return v


def export_xlsx(db: DbcDatabase, out_path: str, include_attrs: bool = True,
                log: Callable[[str], None] = print, group_rows: bool = True,
                message_rows: bool = True, collapse: bool = False) -> dict:
    """group_rows   : plan Excel (Données > Grouper) : les signaux d'un message sont groupés.
    message_rows : ajoute une ligne « parent » par message (comme Plan automatique) ;
                   sinon la 1re ligne de signal sert de ligne parente.
    collapse     : groupes repliés à l'ouverture."""
    cols = _build_columns(db, include_attrs)
    rows, kinds = [], []          # kinds : 'msg' (ligne parente) / 'sig' (ligne enfant)
    for m in db.messages:
        if message_rows and group_rows:
            rows.append([_clean(col.fn(m, None)) for col in cols])
            kinds.append(("msg", m.name))
            for s in m.signals:
                rows.append([_clean(col.fn(m, s)) for col in cols])
                kinds.append(("sig", m.name))
        else:
            for i, s in enumerate(m.signals or [None]):
                rows.append([_clean(col.fn(m, s)) for col in cols])
                kinds.append(("msg" if i == 0 else "sig", m.name))
    if not rows:
        raise ValueError("Aucun message/signal trouvé : ce fichier ne ressemble pas à un DBC valide.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Signaux_DBC"
    ws.append([c.header for c in cols])
    for r in rows:
        ws.append(r)

    n_rows, n_cols = len(rows) + 1, len(cols)
    body_font = Font(name=_FONT, size=10)
    parent_font = Font(name=_FONT, size=10, bold=True)
    parent_fill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
    parent_style = group_rows and message_rows
    hdr_font = Font(name=_FONT, size=10, bold=True, color="FFFFFF")
    widths = [len(c.header) for c in cols]

    # en-tête
    for cell in ws[1]:
        cell.font, cell.fill, cell.border = hdr_font, _HDR_FILL, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    # corps
    for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=n_rows), start=0):
        is_parent = kinds[r_idx][0] == "msg"
        for c_idx, cell in enumerate(row):
            col = cols[c_idx]
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"            # jamais interprété comme une formule
            cell.font, cell.border = body_font, _BORDER
            if is_parent and parent_style:      # ligne « Message » mise en évidence
                cell.font, cell.fill = parent_font, parent_fill
            cell.alignment = Alignment(vertical="top", wrap_text=col.wrap,
                                       horizontal="left" if isinstance(v, str) or v is None else None)
            if col.fmt and isinstance(v, (int, float)):
                cell.number_format = col.fmt
            if v is not None:
                longest = max(len(x) for x in str(v).split("\n"))
                if longest > widths[c_idx]:
                    widths[c_idx] = longest

    for i, col in enumerate(cols):
        ws.column_dimensions[get_column_letter(i + 1)].width = min(max(widths[i] * 1.1 + 3, 10), col.max_width)

    # ---- plan Excel (Outline / Grouper) ------------------------------------
    n_groups = 0
    if group_rows:
        ws.sheet_properties.outlinePr.summaryBelow = False   # « − » sur la ligne parente (au-dessus)
        ws.sheet_properties.outlinePr.summaryRight = True
        ws.sheet_format.outlineLevelRow = 1
        for i, (kind, _name) in enumerate(kinds):
            xl_row = i + 2
            if kind == "sig":
                rd = ws.row_dimensions[xl_row]
                rd.outlineLevel = 1
                rd.hidden = collapse
            elif collapse:
                ws.row_dimensions[xl_row].collapsed = True
        n_groups = sum(1 for k, n in kinds if k == "msg")

    # vrai tableau Excel (les filtres sur toutes les colonnes et les bandes alternées
    # viennent du tableau ; ne PAS ajouter ws.auto_filter en plus : fichier corrompu)
    ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
    tab = Table(displayName="Tableau_DBC", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                        showLastColumn=False, showRowStripes=True,
                                        showColumnStripes=False)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 90

    wb.properties.title = f"Export DBC - {db.source}"
    wb.properties.subject = "Conversion DBC CAN vers Excel"
    wb.properties.description = (
        f"Source : {db.source} | Version DBC : {db.version or 'n/a'} | "
        f"BusType : {db.global_attrs.get('BusType', 'n/a')} | "
        f"DBName : {db.global_attrs.get('DBName', 'n/a')} | "
        f"Généré le {datetime.datetime.now():%Y-%m-%d %H:%M}")
    wb.save(out_path)

    n_sig = sum(len(m.signals) for m in db.messages)
    mux_msgs = [m for m in db.messages if any(s.mux_indicator for s in m.signals)]
    mux_sigs = sum(1 for m in db.messages for s in m.signals if s.mux_indicator)
    return dict(nodes=len(db.nodes), messages=len(db.messages), signals=n_sig,
                rows=len(rows), columns=n_cols, groups=n_groups, mux_messages=len(mux_msgs),
                mux_signals=mux_sigs, attr_columns=len(db.used_attrs) if include_attrs else 0)


def convert(dbc_path: str, xlsx_path: str, include_attrs: bool = True,
            log: Callable[[str], None] = print, group_rows: bool = True,
            message_rows: bool = True, collapse: bool = False) -> dict:
    log(f"Lecture de {dbc_path} ...")
    db = parse_dbc(dbc_path)
    log(f"  {len(db.nodes)} nœuds, {len(db.messages)} messages, "
        f"{sum(len(m.signals) for m in db.messages)} signaux, "
        f"{len(db.attr_defs)} définitions d'attributs.")
    empty = [m.name for m in db.messages if not m.signals]
    if empty:
        log(f"  Attention : {len(empty)} message(s) sans signal (seule la ligne du message est créée).")
    log("Écriture du fichier Excel ...")
    stats = export_xlsx(db, xlsx_path, include_attrs, log, group_rows, message_rows, collapse)
    log(f"  {stats['rows']} lignes x {stats['columns']} colonnes "
        f"(dont {stats['attr_columns']} colonnes d'attributs).")
    if group_rows:
        log(f"  Plan Excel : {stats['groups']} groupes (un par message).")
    log(f"  Multiplexage : {stats['mux_messages']} message(s), {stats['mux_signals']} signal(aux) multiplexé(s)/multiplexeur(s).")
    log(f"Terminé : {xlsx_path}")
    return stats


# =============================================================================
#  INTERFACE GRAPHIQUE (Tkinter, thème moderne)
# =============================================================================

# Palette (assortie à l'en-tête bleu du tableau Excel)
_C_PRIMARY, _C_PRIMARY_HOVER, _C_PRIMARY_DARK = "#1F4E78", "#2B6CA3", "#163A59"
_C_BG, _C_CARD, _C_BORDER = "#F3F5F8", "#FFFFFF", "#D9DEE6"
_C_TEXT, _C_MUTED = "#1E2530", "#6B7585"
_C_OK, _C_ERR = "#1E8E3E", "#C5221F"


def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    # rendu net sur écrans haute densité (Windows)
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)  # type: ignore[attr-defined]
    except Exception:  # noqa: BLE001
        pass

    class App(tk.Tk):
        W = 660

        def __init__(self):
            super().__init__()
            self.title("Convertisseur DBC → Excel")
            self.configure(bg=_C_BG)
            self.resizable(False, False)
            self.q: queue.Queue = queue.Queue()
            self.dbc_var = tk.StringVar()
            self.out_var = tk.StringVar()
            self.attr_var = tk.BooleanVar(value=True)
            self.group_var = tk.BooleanVar(value=True)
            self.msgrow_var = tk.BooleanVar(value=True)
            self.collapse_var = tk.BooleanVar(value=False)
            self.open_var = tk.BooleanVar(value=True)
            self.status_var = tk.StringVar(value="Sélectionnez un fichier .dbc pour commencer.")
            self._auto_out = True
            self._details_open = False
            self._last_out = ""
            self._fonts()
            self._styles()
            self._build()
            self._center()

        # ------------------------------------------------------------- style
        def _fonts(self):
            fam = "Segoe UI" if sys.platform.startswith("win") else "DejaVu Sans"
            self.f_title = (fam, 16, "bold")
            self.f_sub = (fam, 9)
            self.f_base = (fam, 10)
            self.f_bold = (fam, 10, "bold")
            self.f_btn = (fam, 11, "bold")
            self.f_mono = ("Consolas" if sys.platform.startswith("win") else "DejaVu Sans Mono", 9)

        def _styles(self):
            st = ttk.Style(self)
            st.theme_use("clam")
            st.configure(".", font=self.f_base, background=_C_CARD, foreground=_C_TEXT)
            st.configure("Card.TFrame", background=_C_CARD)
            st.configure("Bg.TFrame", background=_C_BG)
            st.configure("TLabel", background=_C_CARD, foreground=_C_TEXT)
            st.configure("Muted.TLabel", foreground=_C_MUTED, background=_C_CARD)
            st.configure("Bg.TLabel", background=_C_BG, foreground=_C_TEXT)
            st.configure("Status.TLabel", background=_C_BG, foreground=_C_MUTED, font=self.f_base)
            st.configure("Link.TLabel", background=_C_BG, foreground=_C_PRIMARY, font=self.f_sub)
            st.configure("TEntry", fieldbackground="#FAFBFC", bordercolor=_C_BORDER,
                         lightcolor=_C_BORDER, darkcolor=_C_BORDER, padding=6)
            st.map("TEntry", bordercolor=[("focus", _C_PRIMARY)], lightcolor=[("focus", _C_PRIMARY)],
                   darkcolor=[("focus", _C_PRIMARY)])
            st.configure("TCheckbutton", background=_C_CARD, focuscolor=_C_CARD)
            st.map("TCheckbutton", background=[("active", _C_CARD)],
                   foreground=[("disabled", "#A0A7B3")])
            st.configure("Soft.TButton", background="#EDF1F6", foreground=_C_PRIMARY, borderwidth=0,
                         padding=(12, 6), font=self.f_bold)
            st.map("Soft.TButton", background=[("active", "#DCE5EF"), ("pressed", "#CBD8E7")])
            st.configure("Accent.TButton", background=_C_PRIMARY, foreground="#FFFFFF", borderwidth=0,
                         padding=(18, 10), font=self.f_btn)
            st.map("Accent.TButton",
                   background=[("disabled", "#9DB3C8"), ("pressed", _C_PRIMARY_DARK),
                               ("active", _C_PRIMARY_HOVER)],
                   foreground=[("disabled", "#FFFFFF")])
            st.configure("Thin.Horizontal.TProgressbar", troughcolor="#E3E8EF", background=_C_PRIMARY,
                         bordercolor=_C_BG, lightcolor=_C_PRIMARY, darkcolor=_C_PRIMARY, thickness=4)

        # ------------------------------------------------------ construction
        def _card(self, parent):
            outer = tk.Frame(parent, bg=_C_BORDER)            # bordure fine 1 px
            inner = ttk.Frame(outer, style="Card.TFrame", padding=(16, 12))
            inner.pack(fill="both", expand=True, padx=1, pady=1)
            return outer, inner

        def _build(self):
            # bandeau d'en-tête
            head = tk.Frame(self, bg=_C_PRIMARY, height=74)
            head.pack(fill="x")
            head.pack_propagate(False)
            tk.Label(head, text="DBC  →  Excel", bg=_C_PRIMARY, fg="white",
                     font=self.f_title).pack(anchor="w", padx=22, pady=(12, 0))
            tk.Label(head, text="Convertit un fichier CAN .dbc en un tableau Excel ",
                     bg=_C_PRIMARY, fg="#BFD3E6", font=self.f_sub).pack(anchor="w", padx=22)

            body = ttk.Frame(self, style="Bg.TFrame", padding=(18, 14, 18, 10))
            body.pack(fill="both", expand=True)

            # carte fichiers
            c1, f = self._card(body)
            c1.pack(fill="x")
            f.columnconfigure(1, weight=1)
            ttk.Label(f, text="Fichier DBC", font=self.f_bold).grid(row=0, column=0, sticky="w")
            ttk.Entry(f, textvariable=self.dbc_var, width=50).grid(row=1, column=0, columnspan=2,
                                                                  sticky="ew", pady=(3, 10), padx=(0, 8))
            ttk.Button(f, text="Parcourir…", style="Soft.TButton",
                       command=self.pick_dbc).grid(row=1, column=2, pady=(3, 10))
            ttk.Label(f, text="Fichier Excel de sortie", font=self.f_bold).grid(row=2, column=0, sticky="w")
            e = ttk.Entry(f, textvariable=self.out_var)
            e.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(3, 0), padx=(0, 8))
            e.bind("<Key>", lambda _e: setattr(self, "_auto_out", False))
            ttk.Button(f, text="Choisir…", style="Soft.TButton",
                       command=self.pick_out).grid(row=3, column=2, pady=(3, 0))

            # carte options
            c2, o = self._card(body)
            c2.pack(fill="x", pady=(10, 0))
            ttk.Label(o, text="Options", font=self.f_bold).pack(anchor="w", pady=(0, 4))
            ttk.Checkbutton(o, text="Une colonne par attribut du DBC",
                            variable=self.attr_var).pack(anchor="w", pady=1)
            ttk.Checkbutton(o, text="Grouper les signaux par message (plan Excel)",
                            variable=self.group_var, command=self._sync_opts).pack(anchor="w", pady=1)
            sub = ttk.Frame(o, style="Card.TFrame")
            sub.pack(anchor="w", padx=(24, 0))
            self.cb_msgrow = ttk.Checkbutton(sub, text="Ligne « Message » parente au-dessus des signaux",
                                             variable=self.msgrow_var)
            self.cb_msgrow.pack(anchor="w", pady=1)
            self.cb_collapse = ttk.Checkbutton(sub, text="Groupes repliés à l'ouverture",
                                               variable=self.collapse_var)
            self.cb_collapse.pack(anchor="w", pady=1)
            ttk.Checkbutton(o, text="Ouvrir Excel à la fin de la conversion",
                            variable=self.open_var).pack(anchor="w", pady=1)

            # action + progression
            act = ttk.Frame(body, style="Bg.TFrame")
            act.pack(fill="x", pady=(12, 0))
            self.btn = ttk.Button(act, text="Convertir", style="Accent.TButton", command=self.start)
            self.btn.pack(side="left")
            self.btn_open = ttk.Button(act, text="Ouvrir le fichier", style="Soft.TButton",
                                       command=self.open_result)
            self.btn_dir = ttk.Button(act, text="Ouvrir le dossier", style="Soft.TButton",
                                      command=self.open_folder)
            self.bar = ttk.Progressbar(body, mode="indeterminate", style="Thin.Horizontal.TProgressbar")
            self.bar.pack(fill="x", pady=(10, 0))

            # statut sur une ligne + lien pour dérouler le journal
            row = ttk.Frame(body, style="Bg.TFrame")
            row.pack(fill="x", pady=(6, 0))
            self.status = ttk.Label(row, textvariable=self.status_var, style="Status.TLabel",
                                    wraplength=self.W - 170)
            self.status.pack(side="left", fill="x", expand=True)
            self.link = ttk.Label(row, text="Détails ▸", style="Link.TLabel", cursor="hand2")
            self.link.pack(side="right")
            self.link.bind("<Button-1>", lambda _e: self.toggle_details())

            # journal compact (masqué par défaut)
            self.log_wrap = tk.Frame(body, bg=_C_BORDER)
            self.log = tk.Text(self.log_wrap, height=5, state="disabled", wrap="word", relief="flat",
                               bg="#FAFBFC", fg=_C_TEXT, font=self.f_mono, padx=8, pady=6, bd=0)
            sb = ttk.Scrollbar(self.log_wrap, orient="vertical", command=self.log.yview)
            self.log.configure(yscrollcommand=sb.set)
            sb.pack(side="right", fill="y", padx=(0, 1), pady=1)
            self.log.pack(side="left", fill="both", expand=True, padx=(1, 0), pady=1)

            self._sync_opts()
            self.update_idletasks()
            self.geometry(f"{self.W}x{self.winfo_reqheight()}")

        def _center(self):
            self.update_idletasks()
            w, h = self.W, self.winfo_reqheight()
            x = (self.winfo_screenwidth() - w) // 2
            y = max((self.winfo_screenheight() - h) // 3, 0)
            self.geometry(f"{w}x{h}+{x}+{y}")

        # ------------------------------------------------------------ actions
        def _sync_opts(self):
            st = "normal" if self.group_var.get() else "disabled"
            self.cb_msgrow.configure(state=st)
            self.cb_collapse.configure(state=st)

        def toggle_details(self):
            self._details_open = not self._details_open
            if self._details_open:
                self.log_wrap.pack(fill="x", pady=(6, 0))
                self.link.configure(text="Détails ▾")
            else:
                self.log_wrap.pack_forget()
                self.link.configure(text="Détails ▸")
            self.update_idletasks()
            self.geometry(f"{self.W}x{self.winfo_reqheight()}")

        def write(self, msg: str):
            self.log.configure(state="normal")
            self.log.insert("end", msg + "\n")
            self.log.see("end")
            self.log.configure(state="disabled")

        def set_status(self, text: str, color: str = _C_MUTED):
            self.status_var.set(text)
            self.status.configure(foreground=color)

        def pick_dbc(self):
            p = filedialog.askopenfilename(title="Sélectionner un fichier DBC",
                                           filetypes=[("Fichiers DBC", "*.dbc"), ("Tous les fichiers", "*.*")])
            if p:
                self.dbc_var.set(p)
                if self._auto_out or not self.out_var.get():
                    self.out_var.set(os.path.splitext(p)[0] + ".xlsx")
                    self._auto_out = True
                self.set_status("Prêt à convertir.")

        def pick_out(self):
            p = filedialog.asksaveasfilename(title="Enregistrer le fichier Excel",
                                             defaultextension=".xlsx",
                                             initialfile=os.path.basename(self.out_var.get() or "export.xlsx"),
                                             filetypes=[("Classeur Excel", "*.xlsx")])
            if p:
                self.out_var.set(p)
                self._auto_out = False

        def open_result(self):
            if self._last_out and hasattr(os, "startfile"):
                os.startfile(self._last_out)  # type: ignore[attr-defined]

        def open_folder(self):
            if self._last_out and hasattr(os, "startfile"):
                os.startfile(os.path.dirname(self._last_out))  # type: ignore[attr-defined]

        def start(self):
            src, dst = self.dbc_var.get().strip(), self.out_var.get().strip()
            if not src or not os.path.isfile(src):
                messagebox.showwarning("Fichier DBC", "Veuillez sélectionner un fichier DBC existant.")
                return
            if not dst:
                dst = os.path.splitext(src)[0] + ".xlsx"
            if not dst.lower().endswith(".xlsx"):
                dst += ".xlsx"
            self.out_var.set(dst)
            self.btn_open.pack_forget()
            self.btn_dir.pack_forget()
            self.btn.configure(state="disabled", text="Conversion…")
            self.bar.configure(mode="indeterminate")
            self.bar.start(12)
            self.set_status("Conversion en cours…")
            self.write("─" * 48)
            threading.Thread(
                target=self._work, daemon=True,
                args=(src, dst, self.attr_var.get(), self.group_var.get(),
                      self.msgrow_var.get(), self.collapse_var.get())).start()
            self.after(100, self._poll)

        def _work(self, src, dst, attrs, group, msgrow, collapse):
            try:
                st = convert(src, dst, attrs, log=lambda m: self.q.put(("log", m)),
                             group_rows=group, message_rows=msgrow, collapse=collapse)
                self.q.put(("ok", (dst, st)))
            except PermissionError:
                self.q.put(("err", "Écriture impossible : le fichier est sans doute ouvert dans Excel. "
                                   "Fermez-le ou changez de nom."))
            except Exception as exc:  # noqa: BLE001
                self.q.put(("err", f"{type(exc).__name__} : {exc}"))

        def _poll(self):
            done = False
            try:
                while True:
                    kind, payload = self.q.get_nowait()
                    if kind == "log":
                        self.write(payload)
                    elif kind == "ok":
                        done = True
                        dst, st = payload
                        self._last_out = dst
                        self.set_status(f"✔ Terminé : {st['messages']} messages, {st['signals']} signaux, "
                                        f"{st['columns']} colonnes.", _C_OK)
                        self.btn_open.pack(side="left", padx=(10, 0))
                        self.btn_dir.pack(side="left", padx=(6, 0))
                        if self.open_var.get():
                            self.open_result()
                    elif kind == "err":
                        done = True
                        self.write("Erreur : " + payload)
                        self.set_status("✘ " + payload, _C_ERR)
                        if not self._details_open:
                            self.toggle_details()
            except queue.Empty:
                pass
            if done:
                self.bar.stop()
                self.bar.configure(mode="determinate", value=0)
                self.btn.configure(state="normal", text="Convertir")
            else:
                self.after(100, self._poll)

    App().mainloop()


if __name__ == "__main__":
    if len(sys.argv) >= 2:
        src_ = sys.argv[1]
        dst_ = sys.argv[2] if len(sys.argv) >= 3 else os.path.splitext(src_)[0] + ".xlsx"
        convert(src_, dst_)
    else:
        run_gui()